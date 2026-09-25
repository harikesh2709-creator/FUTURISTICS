"""
Build Official SIH 2026 Ultimate Presentation Deck matching the video's layout.
"""

import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import win32com.client

SRC_TEMPLATE = r"C:\Users\Harik\Downloads\SIH2026-IDEA-Presentation-Format.pptx"
OUT_DIR = r"c:\vs studio\freight-forecast"
OUT_PPTX_7 = os.path.join(OUT_DIR, "SIH2026_FreightForecast_Pro_Ultimate.pptx")
OUT_PPTX_6 = os.path.join(OUT_DIR, "SIH2026_FreightForecast_Pro_Ultimate_6Slides.pptx")
OUT_PDF_6 = os.path.join(OUT_DIR, "SIH2026_FreightForecast_Pro_Ultimate_6Slides.pdf")
OUT_PDF_7 = os.path.join(OUT_DIR, "SIH2026_FreightForecast_Pro_Ultimate.pdf")

COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_DARK_NAVY = RGBColor(15, 23, 42)
COLOR_TITLE_BLUE = RGBColor(30, 58, 138)

FONT_FAMILY = "Arial"

def style_run(run, text, font_name=FONT_FAMILY, size_pt=12, bold=False, italic=False, color=COLOR_DARK_NAVY):
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color

def update_header_and_footer(slide, slide_num, title_text):
    """Update slide title, top-left team oval, and bottom footer."""
    for shp in slide.shapes:
        # Top-left oval / badge
        if "Oval" in shp.name or (shp.has_text_frame and "Your Team" in shp.text):
            shp.left = Inches(0.40)
            shp.top = Inches(0.25)
            shp.width = Inches(2.20)
            shp.height = Inches(0.72)
            shp.fill.solid()
            shp.fill.fore_color.rgb = COLOR_DARK_NAVY
            shp.line.color.rgb = COLOR_TITLE_BLUE
            shp.line.width = Pt(2.0)
            tf = shp.text_frame
            tf.clear()
            tf.word_wrap = False
            tf.margin_top = Inches(0.05)
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            style_run(p.add_run(), "FUTURISTICS\n", size_pt=11, bold=True, color=COLOR_WHITE)
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.CENTER
            style_run(p2.add_run(), "TEAM ID: SIH-2026", size_pt=8, bold=True, color=RGBColor(56, 189, 248))

        # Slide title
        if "Title" in shp.name or (shp.has_text_frame and ("IDEA TITLE" in shp.text or "TECHNICAL" in shp.text or "FEASIBILITY" in shp.text or "IMPACT" in shp.text or "RESEARCH" in shp.text)):
            shp.left = Inches(2.80)
            shp.top = Inches(0.10)
            shp.width = Inches(9.00)
            shp.height = Inches(0.50)
            tf = shp.text_frame
            tf.clear()
            
            p1 = tf.paragraphs[0]
            p1.alignment = PP_ALIGN.LEFT
            style_run(p1.add_run(), "SMART INDIA HACKATHON 2026", size_pt=10, bold=True, color=COLOR_TITLE_BLUE)
            
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.LEFT
            style_run(p2.add_run(), title_text, size_pt=20, bold=True, color=COLOR_DARK_NAVY)

        # Footer text - removing default footers since our infographics already have them or they are not needed.
        if "Footer" in shp.name or (shp.has_text_frame and "@SIH" in shp.text):
            sp = shp._element
            sp.getparent().remove(sp)

        # Slide Number
        if "Slide Number" in shp.name or (shp.has_text_frame and shp.text.strip().isdigit()):
            sp = shp._element
            sp.getparent().remove(sp)

def clear_old_body_shapes(slide):
    to_remove = []
    for shp in slide.shapes:
        if shp.has_text_frame:
            txt = shp.text_frame.text.lower()
            if "proposed solution" in txt or "technologies to be used" in txt or "analysis of the feasibility" in txt or "potential impact" in txt or "details / links" in txt or "2-3 lines describing" in txt or "describe your idea" in txt:
                to_remove.append(shp)
    for shp in to_remove:
        sp = shp._element
        sp.getparent().remove(sp)

def embed_graphics(slide, image_path):
    if os.path.exists(image_path):
        left = Inches(0.40)
        top = Inches(1.10)
        width = Inches(12.53)
        height = Inches(6.0)
        slide.shapes.add_picture(image_path, left, top, width, height)

def build_all_slides():
    print(f"Loading base template: {SRC_TEMPLATE}")
    prs = Presentation(SRC_TEMPLATE)

    # =========================================================================
    # SLIDE 1: TITLE PAGE
    # =========================================================================
    print("Populating Slide 1: TITLE PAGE...")
    slide1 = prs.slides[0]
    
    for shp in slide1.shapes:
        if "Oval" in shp.name or (shp.has_text_frame and "Your Team" in shp.text):
            shp.left = Inches(0.80)
            shp.top = Inches(0.40)
            shp.width = Inches(2.40)
            shp.height = Inches(0.80)
            shp.fill.solid()
            shp.fill.fore_color.rgb = COLOR_DARK_NAVY
            shp.line.color.rgb = COLOR_TITLE_BLUE
            shp.line.width = Pt(2.0)
            tf = shp.text_frame
            tf.clear()
            tf.word_wrap = False
            tf.margin_top = Inches(0.08)
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            style_run(p.add_run(), "FUTURISTICS\n", size_pt=12, bold=True, color=COLOR_WHITE)
            p2 = tf.add_paragraph()
            p2.alignment = PP_ALIGN.CENTER
            style_run(p2.add_run(), "TEAM ID: SIH-2026", size_pt=9, bold=True, color=RGBColor(56, 189, 248))

        if shp.has_text_frame and ("Problem Statement ID" in shp.text or "Problem Statement Title" in shp.text):
            tf = shp.text_frame
            tf.clear()
            tf.margin_left = Inches(0.2)
            tf.margin_top = Inches(0.2)
            
            meta_items = [
                ("Problem Statement ID", "SIH26006 / SIH2026-LOG-01"),
                ("Problem Statement Title", "FreightForecast Pro : AI Freight Forecasting & Vessel Chartering Optimizer"),
                ("Theme", "Smart Logistics / Maritime & Port Supply Chain"),
                ("PS Category", "Software"),
                ("Team ID", "[As Registered on SIH Portal]"),
                ("Team Name", "FUTURISTICS")
            ]
            
            for idx, (lbl, val) in enumerate(meta_items):
                p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                p.space_after = Pt(25)
                r_lbl = p.add_run()
                style_run(r_lbl, f"• {lbl} – ", size_pt=20, bold=True, color=COLOR_DARK_NAVY)
                r_val = p.add_run()
                style_run(r_val, val, size_pt=20, bold=(lbl in ["Problem Statement ID", "Team Name", "PS Category"]), color=COLOR_TITLE_BLUE)
                
                if lbl == "Team Name":
                    r_val.font.color.rgb = RGBColor(22, 163, 74)

        if shp.has_text_frame and "SMART INDIA HACKATHON 2026" in shp.text:
            tf = shp.text_frame
            for p in tf.paragraphs:
                p.alignment = PP_ALIGN.LEFT
                for r in p.runs:
                    style_run(r, r.text, size_pt=28, bold=True, color=COLOR_TITLE_BLUE)
            shp.left = Inches(3.4)
            shp.top = Inches(0.6)

        if shp.has_text_frame and "TITLE PAGE" in shp.text:
            sp = shp._element
            sp.getparent().remove(sp)

    # =========================================================================
    # SLIDE 2 to 7
    # =========================================================================
    slides_data = [
        (1, "FreightForecast Pro: Predictive Chartering & Route Optimizer", r"c:\vs studio\freight-forecast\sih_slide2_body.png"),
        (2, "TECHNICAL APPROACH & SYSTEM ARCHITECTURE", r"c:\vs studio\freight-forecast\sih_slide3_body.png"),
        (3, "FEASIBILITY, RISK ANALYSIS & MITIGATION MATRIX", r"c:\vs studio\freight-forecast\sih_slide4_body.png"),
        (4, "IMPACT, QUANTIFIABLE ROI & ECOSYSTEM BENEFITS", r"c:\vs studio\freight-forecast\sih_slide5_body.png"),
        (5, "RESEARCH, INDUSTRY STANDARDS & MARKET VALIDATION", r"c:\vs studio\freight-forecast\sih_slide6_body.png"),
        (6, "LIVE PROTOTYPE SHOWCASE & PRODUCTION TECH STACK", r"c:\vs studio\freight-forecast\sih_slide7_body.png")
    ]
    
    for slide_idx, title, img_path in slides_data:
        print(f"Populating Slide {slide_idx+1}...")
        slide = prs.slides[slide_idx]
        update_header_and_footer(slide, slide_idx+1, title)
        clear_old_body_shapes(slide)
        embed_graphics(slide, img_path)

    # Save 7-slide version
    prs.save(OUT_PPTX_7)
    print(f"Saved 7-slide presentation: {OUT_PPTX_7}")

    # Create 6-slide submission version
    rId = prs.slides._sldIdLst[6].rId
    prs.part.drop_rel(rId)
    del prs.slides._sldIdLst[6]
    prs.save(OUT_PPTX_6)
    print(f"Saved 6-slide submission presentation: {OUT_PPTX_6}")

    print("Exporting presentations to PDF via PowerPoint COM...")
    try:
        ppt_app = win32com.client.Dispatch("PowerPoint.Application")
        ppt_app.Visible = 1
        
        pres6 = ppt_app.Presentations.Open(OUT_PPTX_6, WithWindow=False)
        pres6.SaveAs(OUT_PDF_6, 32)
        pres6.Close()
        print(f"Exported PDF (6 slides): {OUT_PDF_6}")

        pres7 = ppt_app.Presentations.Open(OUT_PPTX_7, WithWindow=False)
        pres7.SaveAs(OUT_PDF_7, 32)
        pres7.Close()
        print(f"Exported PDF (7 slides): {OUT_PDF_7}")

        ppt_app.Quit()
    except Exception as e:
        print("Warning during PDF export:", e)
        
if __name__ == "__main__":
    build_all_slides()
