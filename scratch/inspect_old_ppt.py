from pptx import Presentation

prs = Presentation('SPECTRA_SIH_2026_Presentation_Previous_AI_Copy.pptx')
print(f"Total slides: {len(prs.slides)}")
for idx, slide in enumerate(prs.slides, 1):
    texts = []
    images = []
    for shape in slide.shapes:
        if shape.shape_type == 13: # Picture
            images.append(shape.name)
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                txt = p.text.strip()
                if txt:
                    texts.append((round(p.font.size.pt, 1) if p.font and p.font.size else None, txt))
    print(f"\n--- SLIDE {idx} (Shapes: {len(slide.shapes)}, Images: {len(images)}) ---")
    for size, t in texts[:8]:
        print(f"  [{size} pt] {t[:80]}")
