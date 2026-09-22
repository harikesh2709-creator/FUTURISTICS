from pptx import Presentation

prs = Presentation('SPECTRA_SIH_2026_Presentation_Raw_Original_Backup.pptx')
print(f"Total slides: {len(prs.slides)}")

for i, slide in enumerate(prs.slides, 1):
    print(f"\n==================== SLIDE {i} ====================")
    for s_idx, shape in enumerate(slide.shapes):
        if shape.shape_type == 13: # Picture
            # Try to get image filename or part name if available
            try:
                pname = shape.image.part.partname
            except Exception:
                pname = "unknown"
            print(f"  [PICTURE] Shape {s_idx}: name='{shape.name}', pos=({shape.left.inches:.2f}, {shape.top.inches:.2f}, w={shape.width.inches:.2f}, h={shape.height.inches:.2f}), part={pname}")
        elif shape.has_text_frame:
            txt = shape.text_frame.text.replace('\n', ' ')
            if txt.strip():
                sample = txt.strip()[:60].encode('ascii', 'replace').decode('ascii')
                sizes = [round(p.font.size.pt, 1) for p in shape.text_frame.paragraphs if p.font and p.font.size]
                runs_sizes = [round(r.font.size.pt, 1) for p in shape.text_frame.paragraphs for r in p.runs if r.font and r.font.size]
                all_sizes = sorted(list(set(sizes + runs_sizes)))
                print(f"  [TEXT] Shape {s_idx}: pos=({shape.left.inches:.2f}, {shape.top.inches:.2f}, w={shape.width.inches:.2f}, h={shape.height.inches:.2f}) sizes={all_sizes} | '{sample}'")
