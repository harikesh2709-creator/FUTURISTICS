from pptx import Presentation

prs = Presentation(r'C:\Users\Harik\Downloads\SIH2026-Idea-Presentation_Final.pptx')
print(f'Total slides: {len(prs.slides)}')
for idx, slide in enumerate(prs.slides):
    print(f'=== Slide {idx+1} ===')
    for s in slide.shapes:
        text = ''
        if s.has_text_frame and s.text_frame.text:
            text = s.text_frame.text.replace('\n', ' ')[:70]
        print(f'  Shape: {s.name:<30} | Type: {s.shape_type} | L: {s.left/914400:0.2f}\", T: {s.top/914400:0.2f}\", W: {s.width/914400:0.2f}\", H: {s.height/914400:0.2f}\" | Text: {text}')
