from pptx import Presentation

prs = Presentation('SPECTRA_SIH_2026_Presentation_Original_8Slide.pptx')
for idx, slide in enumerate(prs.slides, 1):
    print(f"\n==================== SLIDE {idx} ====================")
    for s_idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            for p_idx, p in enumerate(shape.text_frame.paragraphs):
                txt = p.text.strip().replace('\n', ' ')
                if txt:
                    p_size = round(p.font.size.pt, 1) if p.font and p.font.size else None
                    run_sizes = [round(r.font.size.pt, 1) for r in p.runs if r.font and r.font.size]
                    effective_size = run_sizes[0] if run_sizes else p_size
                    print(f"  Shape {s_idx} P{p_idx}: size={effective_size} (p_size={p_size}, runs={run_sizes}) | text='{txt[:50]}'")
        elif shape.has_table:
            for r_idx, row in enumerate(shape.table.rows):
                for c_idx, cell in enumerate(row.cells):
                    for p in cell.text_frame.paragraphs:
                        txt = p.text.strip().replace('\n', ' ')
                        if txt:
                            p_size = round(p.font.size.pt, 1) if p.font and p.font.size else None
                            run_sizes = [round(r.font.size.pt, 1) for r in p.runs if r.font and r.font.size]
                            effective_size = run_sizes[0] if run_sizes else p_size
                            print(f"  Table R{r_idx}C{c_idx}: size={effective_size} | text='{txt[:40]}'")
