from pptx import Presentation
import os

pptx_path = r'c:\vs studio\ntro-signal-analyzer\SPECTRA_SIH_2026_Presentation.pptx'
prs = Presentation(pptx_path)
print(f"Presentation loaded: {len(prs.slides)} slides.")

for i, slide in enumerate(prs.slides):
    text_count = 0
    img_count = 0
    for shape in slide.shapes:
        if shape.has_text_frame:
            text_count += 1
        if shape.shape_type == 13: # picture
            img_count += 1
    print(f"Slide {i+1}: {len(slide.shapes)} shapes, {text_count} text frames, {img_count} images")
