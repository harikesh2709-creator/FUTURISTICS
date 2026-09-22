import pptx, os

prs = pptx.Presentation('SIH2026_FreightForecast_Pro_2.pptx')
for s_idx, slide in enumerate(prs.slides, 1):
    print(f'=== SLIDE {s_idx} ===')
    for shp in slide.shapes:
        if shp.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            blip = shp._element.xpath('.//a:blip')
            rId = blip[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed') if blip else None
            target = slide.part.rels[rId].target_ref if rId else 'unknown'
            print(f'  Picture shape "{shp.name}" -> {target} | left={shp.left.inches:.2f}, top={shp.top.inches:.2f}, w={shp.width.inches:.2f}, h={shp.height.inches:.2f}')
